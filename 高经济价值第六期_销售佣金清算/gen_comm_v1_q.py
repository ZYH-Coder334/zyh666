#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""高经济价值第六期·销售佣金与渠道返利历史清算 v1 —— 题目数据生成器"""
import random, os
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

random.seed(20261023)
BASE = '/workspace/高经济价值第六期_销售佣金清算'
os.makedirs(BASE, exist_ok=True)

REF = date(2025, 8, 31)
HDR_FILL = PatternFill('solid', fgColor='2F5597')
HDR_FONT = Font(name='宋体', bold=True, color='FFFFFF', size=10)
BODY = Font(name='宋体', size=9)
THIN = Side(style='thin', color='CCCCCC')
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        w = (widths or {}).get(c, 14)
        ws.column_dimensions[get_column_letter(c)].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.font = BODY
            cell.alignment = WRAP
            cell.border = BORDER
    ws.freeze_panes = 'A2'
    return ws


def dstr(d):
    return d.strftime('%Y-%m-%d')


# ---------------- 基础主数据 ----------------
TEAMS = [
    ('华东', '江苏、浙江、上海、安徽', 11500),
    ('华北', '北京、天津、河北、山东', 8000),
    ('华南', '广东、福建、广西、海南', 6800),
    ('西南', '四川、重庆、云南、贵州', 7000),
    ('西北', '陕西、甘肃、宁夏、新疆', 7300),
    ('华中', '湖北、湖南、河南、江西', 6700),
]
SURNAMES = list('王李张刘陈杨黄赵吴周徐孙马朱胡郭何高林罗郑梁谢宋唐许韩冯邓曹彭曾肖田董袁潘于蒋蔡余杜叶程苏魏吕丁任沈姚卢姜崔钟谭陆汪范金石廖贾夏韦付方白邹孟熊秦邱江尹薛段雷侯龙史陶黎贺顾毛郝龚邵万钱严覃武戴莫孔向汤')
GIVENS = ['伟', '芳', '娜', '敏', '静', '丽', '强', '磊', '军', '洋', '勇', '艳', '杰', '娟', '涛', '明', '超', '霞', '平', '刚', '建国', '志强', '文博', '雨欣', '浩然', '子涵', '嘉怡', '思远', '晨曦', '若曦', '泽宇', '欣怡', '梓萱', '宇航', '梦琪', '俊杰', '诗涵', '天佑', '一鸣', '嘉豪']
TITLES = ['销售总监', '大区经理', '客户经理', '客户经理', '客户经理', '客户经理', '客户经理', '客户经理', '客户经理', '客户经理',
          '客户经理', '客户经理', '客户经理', '客户经理', '客户经理', '客户经理', '客户经理', '渠道经理', '渠道经理', '客户经理']
SALARY = {'销售总监': 30000, '大区经理': 22000, '渠道经理': 15000, '客户经理': 13000}
MTARGET = {'销售总监': 80, '大区经理': 55, '渠道经理': 35, '客户经理': 32}

names = []
used = set()
for i in range(120):
    while True:
        nm = random.choice(SURNAMES) + random.choice(GIVENS)
        if nm not in used:
            used.add(nm)
            break
    names.append(nm)

roster = []
for ti, (team, area, tgt24) in enumerate(TEAMS):
    for j in range(20):
        idx = ti * 20 + j
        eid = f'XS{idx + 1:03d}'
        title = TITLES[j]
        if j == 0:
            hire = date(random.randint(2015, 2019), random.randint(1, 12), random.randint(1, 27))
        elif j <= 2:
            hire = date(random.randint(2016, 2020), random.randint(1, 12), random.randint(1, 27))
        else:
            hire = date(random.randint(2017, 2023), random.randint(1, 12), random.randint(1, 27))
        leave = None
        # 每团队 2 名离职
        if j in (5, 14) and ti % 3 == 0:
            leave = date(2023, random.randint(6, 12), random.randint(1, 27))
        elif j in (7, 16) and ti % 3 == 1:
            leave = date(2024, random.randint(3, 11), random.randint(1, 27))
        elif j in (9, 17) and ti % 3 == 2:
            leave = date(2025, random.randint(1, 5), random.randint(1, 27))
        roster.append([eid, names[idx], team, area, title, dstr(hire), dstr(leave) if leave else '', '', 0, 0, ''])

# 10 名 2024-2025 新入职（替换部分 2023 年前入职的客户经理，保留离职标记）
new_idx = [r for r in range(120) if roster[r][4] == '客户经理' and roster[r][6] == ''][:10]
new_dates = [date(2024, 2, 15), date(2024, 4, 8), date(2024, 6, 20), date(2024, 9, 5), date(2024, 11, 12),
             date(2025, 1, 6), date(2025, 2, 18), date(2025, 3, 25), date(2025, 4, 9), date(2025, 6, 16)]
for k, r in enumerate(new_idx):
    roster[r][5] = dstr(new_dates[k])
    roster[r][10] = '2024年后新入职'

for r in roster:
    r[8] = SALARY[r[4]]
    r[9] = MTARGET[r[4]] + (random.randint(-8, 12) if r[4] == '客户经理' else 0)

team_meta = []
for ti, (team, area, tgt24) in enumerate(TEAMS):
    leader = roster[ti * 20][1]
    team_meta.append([team, int(tgt24 * 0.85), int(tgt24 * 0.93), tgt24, int(tgt24 * 1.05), leader,
                      '年度目标按含税回款口径'])

# ---------------- 客户与渠道 ----------------
CITY = ['苏州', '杭州', '上海', '南京', '无锡', '北京', '天津', '济南', '青岛', '广州', '深圳', '东莞',
        '佛山', '成都', '重庆', '昆明', '西安', '兰州', '武汉', '长沙', '郑州', '合肥', '福州', '厦门']
IND = ['精密机械', '重工装备', '化工集团', '电力工程', '水务集团', '新能源汽车', '电子科技', '新材料',
       '造纸集团', '食品机械', '矿山设备', '环保工程', '船舶重工', '轨道交通', '制药设备']
customers = []
for i in range(180):
    cid = f'KH{i + 1:03d}'
    customers.append([cid, f'{random.choice(CITY)}{random.choice(IND)}有限公司', random.choice(CITY),
                      random.choice(['战略客户', '重点客户', '普通客户'])])

CH_TEAM = ['华东', '华北', '华南', '西南', '西北', '华中']
channels = []
for i in range(40):
    cid = f'QD{i + 1:03d}'
    big = i < 5
    channels.append([cid, f'{random.choice(CITY)}联创机电设备有限公司', CH_TEAM[i % 6],
                     dstr(date(random.randint(2018, 2022), random.randint(1, 12), 1)),
                     dstr(date(2026, 12, 31)) if i not in (7, 19, 31) else dstr(date(2025, 6, 30)),
                     300000 if big else 0, roster[(i % 6) * 20 + 17][1],
                     '年度保底返利30万元' if big else '无保底'])

# 渠道返利阶梯
TIERS = [('500万元以下', 0, 500, 0.02), ('500万元（含）至1000万元', 500, 1000, 0.03),
         ('1000万元（含）至2000万元', 1000, 2000, 0.04), ('2000万元（含）以上', 2000, 99999, 0.05)]

CH_WEIGHT = [0.06] * 5 + [0.03] * 10 + [0.016] * 25
CH_WEIGHT = [w / sum(CH_WEIGHT) for w in CH_WEIGHT]

# ---------------- 订单生成 ----------------
orders = []
seq = 0
year_counts = {2022: 170, 2023: 200, 2024: 230, 2025: 120}
team_w = [t[2] for t in TEAMS]
team_w = [w / sum(team_w) for w in team_w]

customer_first = {}
all_orders = []


def pick_sales(team, ship, want_channel=False):
    cand = []
    for r in roster:
        if r[2] != team:
            continue
        if r[5] and date.fromisoformat(r[5]) > ship - timedelta(days=30):
            continue
        if r[6] and date.fromisoformat(r[6]) < ship:
            continue
        if want_channel and r[4] == '渠道经理':
            cand += [(r, 6)]
        elif not want_channel and r[4] in ('客户经理', '大区经理', '销售总监'):
            cand += [(r, {'客户经理': 7, '大区经理': 2, '销售总监': 1}[r[4]])]
        elif not want_channel and r[4] == '渠道经理':
            cand += [(r, 1)]
    if not cand:
        for r in roster:
            if r[2] == team and (not r[6] or date.fromisoformat(r[6]) >= ship):
                cand.append((r, 1))
    total = sum(w for _, w in cand)
    x = random.uniform(0, total)
    acc = 0
    for r, w in cand:
        acc += w
        if x <= acc:
            return r
    return cand[-1][0]


for y, cnt in year_counts.items():
    for k in range(cnt):
        seq += 1
        if y == 2025:
            od = date(2025, random.randint(1, 8), random.randint(1, 20))
        else:
            od = date(y, random.randint(1, 12), random.randint(1, 27))
        ship = od + timedelta(days=random.randint(3, 20))
        if ship > REF:
            ship = REF
        team = random.choices([t[0] for t in TEAMS], weights=team_w)[0]
        is_ch = random.random() < 0.25
        sales = pick_sales(team, ship, want_channel=is_ch)
        if is_ch and sales[4] != '渠道经理' and random.random() < 0.7:
            sales = pick_sales(team, ship, want_channel=True)
        x = random.random()
        if x < 0.55:
            amt_w = random.randint(60, 250)
        elif x < 0.80:
            amt_w = random.randint(250, 500)
        elif x < 0.93:
            amt_w = random.randint(500, 900)
        else:
            amt_w = random.randint(900, 1600)
        if is_ch:
            amt_w = int(amt_w * random.uniform(0.5, 0.9))
        amt = amt_w * 10000
        cust = random.choice(customers)
        ctype = '续单'
        if cust[0] not in customer_first:
            customer_first[cust[0]] = True
            if random.random() < 0.7:
                ctype = '新客户首单'
        ch_id = ''
        if is_ch:
            ch = random.choices(channels, weights=CH_WEIGHT)[0]
            ch_id = ch[0]
            ctype = '渠道单'
        coop = ''
        all_orders.append({'oid': f'DD{y}{seq:04d}', 'od': od, 'ship': ship, 'cust': cust,
                           'ctype': ctype, 'ch': ch_id, 'sales': sales, 'coop': coop, 'amt': amt,
                           'team': team})

# 协作单：2023 年 12 单，2024-2025 年 18 单
coop_pool_23 = [o for o in all_orders if o['ship'].year == 2023 and not o['ch']]
coop_pool_24 = [o for o in all_orders if o['ship'].year in (2024, 2025) and not o['ch']]
for pool, n in ((coop_pool_23, 12), (coop_pool_24, 18)):
    for o in random.sample(pool, n):
        other = [r for r in roster if r[2] != o['team'] and (not r[6] or date.fromisoformat(r[6]) >= o['ship'])]
        o['coop'] = random.choice(other)[0]

# 跨年发货：订单日期 2022 年、发货 2023 年（用于口径陷阱）
cross = [o for o in all_orders if o['ship'].year == 2023 and o['od'].year == 2023][:8]
for k, o in enumerate(cross):
    o['od'] = date(2022, 12, 20 + k)

# 重复导入：12 单
dups = random.sample(all_orders, 12)
order_rows = []
for o in all_orders:
    order_rows.append([o['oid'], dstr(o['od']), dstr(o['ship']), o['cust'][0], o['cust'][1], o['cust'][3],
                       o['ctype'], o['ch'], o['sales'][0], o['coop'], o['amt'], 0.13, o['team'], ''])
for o in dups:
    order_rows.append([o['oid'], dstr(o['od']), dstr(o['ship']), o['cust'][0], o['cust'][1], o['cust'][3],
                       o['ctype'], o['ch'], o['sales'][0], o['coop'], int(o['amt'] * 1.02), 0.13, o['team'],
                       '系统迁移重复导入'])

# ---------------- 回款 / 退货 / 坏账 ----------------
pay_rows = []
returns = []
bad = []
pseq = 0
for o in all_orders:
    r = random.random()
    if r < 0.68:
        mode = 'full'
    elif r < 0.82:
        mode = 'high'
    elif r < 0.95:
        mode = 'low'
    else:
        mode = 'none'
    first = int(o['amt'] * 0.3)
    if mode == 'full':
        rest = o['amt'] - first
    elif mode == 'high':
        rest = int((o['amt'] - first) * random.uniform(0.86, 0.985))
    elif mode == 'low':
        rest = int((o['amt'] - first) * random.uniform(0.4, 0.83))
    else:
        rest = 0
    pays = [(o['ship'] + timedelta(days=random.randint(10, 30)), first)]
    if rest > 0:
        pays.append((o['ship'] + timedelta(days=random.randint(55, 150)), rest))
    for pd, pa in pays:
        if pd > REF:
            continue
        pseq += 1
        pay_rows.append([f'HK{pseq:05d}', o['oid'], dstr(pd), pa, '首款' if pa == first else '尾款',
                         f'BL{random.randint(100000, 999999)}', ''])
    # 退货
    if random.random() < 0.10 and o['ship'] <= REF - timedelta(days=40):
        rd = o['ship'] + timedelta(days=random.randint(40, 170))
        if rd <= REF:
            ra = int(o['amt'] * random.choice([0.1, 0.2, 0.3, 0.5, 0.8]))
            returns.append([f'TH{len(returns) + 1:04d}', o['oid'], dstr(rd), ra,
                            random.choice(['型号选错', '质量异议', '项目取消', '重复下单']), '已批准'])
    # 坏账
    if mode in ('none', 'low') and random.random() < 0.36:
        paid = sum(p for _, p in pays if _ <= REF)
        bad.append([f'HZ{len(bad) + 1:04d}', o['oid'], o['cust'][1], dstr(date(2025, 6, 30)),
                    o['amt'] - paid, '逾期超过180天，经法务确认无法收回'])

# ---------------- 渠道返利已付 ----------------
rebate_paid = []
for ch in channels:
    for y in (2022, 2023, 2024):
        net = sum(o['amt'] for o in all_orders if o['ch'] == ch[0] and o['ship'].year == y)
        paid = int(net * random.choice([0.02, 0.025, 0.03]) * random.uniform(0.7, 1.05)) if net else 0
        rebate_paid.append([ch[0], ch[1], y, paid, '年度预结', ''])

# ---------------- 已发佣金台账 ----------------
paid_rows = []
for r in roster:
    for y in (2022, 2023, 2024, 2025):
        base = sum(o['amt'] for o in all_orders if o['sales'][0] == r[0] and o['ship'].year == y and o['ship'] <= REF)
        paid = int(base * 0.03 * 0.7)
        note = ''
        if r[0] in ('XS006', 'XS032', 'XS055', 'XS078', 'XS099', 'XS114'):
            paid = int(paid * 1.08)
            note = '按含税开票额预发'
        elif r[0] in ('XS021', 'XS044', 'XS067', 'XS090'):
            paid = int(paid * 0.85)
            note = '预发比例错误'
        elif r[0] in ('XS013', 'XS058', 'XS103'):
            paid = 0
            note = '待核实未发放'
        paid_rows.append([r[0], r[1], y, paid, f'{y}年度第{random.randint(1, 4)}批', note])

# ---------------- 争议 ----------------
DISPUTE_REASON = [
    '按合同额×3%计算，公司少发', '本人已离职，但订单回款发生在离职后，主张全额',
    '渠道单应按3%而不是1.5%', '合作单应全部归本人，不应拆分',
    '退货不应冲减佣金', '该订单已回款，公司以未验收为由拒发',
    '坏账核销不应扣减佣金', '新人保底应按全年发放',
    '系统重复订单重复计提被公司核减，主张恢复', 'V2应按4月1日起执行，1-3月仍按V1计算',
    '封顶条款未在合同中约定，主张取消', '团队池应按人头平均分配']
disputes = []
dcand = [o for o in all_orders if o['amt'] >= 3000000]
for i, o in enumerate(random.sample(dcand, 60), 1):
    claim = int(o['amt'] * random.choice([0.03, 0.035, 0.04, 0.05]))
    disputes.append([f'ZY{i:03d}', o['sales'][0], o['sales'][1], o['oid'], dstr(date(2025, random.randint(6, 8), random.randint(1, 28))),
                     claim, random.choice(DISPUTE_REASON),
                     random.choice(['订单截图', '回款流水', '离职证明', '合同扫描件', '微信沟通记录'])])

# ---------------- 规则文本 ----------------
V1 = [['V1-1', '适用期间', '2022-01-01至2022-12-31，按订单发货日期归属'],
      ['V1-2', '计提基数', '订单含税发货额扣除退货金额'],
      ['V1-3', '计提费率', '标准3%；新客户首单4%'],
      ['V1-4', '发放条件', '订单累计回款率≥90%（（含税回款-退货）/含税发货额）'],
      ['V1-5', '未达条件', '挂账不发放，回款达标后补发'],
      ['V1-6', '月度封顶', '个人当月应发不超过10倍月基本工资，超出部分不结转'],
      ['V1-7', '合作单', '全部归签约销售'],
      ['V1-8', '离职', '离职当月已发货订单照常计提']]
V2 = [['V2-1', '适用期间', '2023-01-01至2023-12-31，按订单发货日期归属'],
      ['V2-2', '计提基数', '订单含税回款额扣除退款'],
      ['V2-3', '计提费率', '3.5%；单笔订单累计含税回款超过500万元的部分按5%'],
      ['V2-4', '封顶', '无封顶'],
      ['V2-5', '合作单', '全部归签约销售'],
      ['V2-6', '离职', '离职当月计提，次月起回款不再计提'],
      ['V2-7', '退款', '按退款金额×3.5%冲减']]
V3 = [['V3-1', '适用期间', '2024-01-01至2025-08-31，按订单发货日期归属'],
      ['V3-2', '计提基数', '不含税净回款＝（含税回款-退款-坏账核销）/1.13'],
      ['V3-3', '计提费率', '直销3%；渠道单1.5%'],
      ['V3-4', '分配结构', '个人80%，团队池20%'],
      ['V3-5', '合作单', '主谈70%，协助30%'],
      ['V3-6', '月度封顶', '个人当月应发不超过12倍月基本工资，超出部分不结转'],
      ['V3-7', '离职', '离职次月起回款不再计提'],
      ['V3-8', '新人保护', '入职前3个自然月，月佣金不足8000元补足至8000元'],
      ['V3-9', '团队池', '团队年度目标达成率≥100%方可发放，未达标留存']]
NOTICES = [['睿联销字〔2023〕07号', '2023-03-20', '2023-01-01', '全体销售', '发布《佣金政策V2》，自2023年1月1日起追溯执行，替代V1'],
           ['睿联销字〔2024〕02号', '2023-12-28', '2024-01-01', '全体销售', '发布《佣金政策V3》，调整计提基数为不含税净回款，引入团队池'],
           ['睿联销字〔2024〕15号', '2024-07-01', '2024-07-01', '渠道销售', '明确渠道单佣金按直销费率减半执行，避免与渠道返利重复']]
MEETINGS = [['2023-03-10', '一季度经营分析会', '总经理、销售VP、财务总监、各大区负责人',
             '讨论佣金政策调整，倾向自2023-04-01执行且不追溯', '讨论稿，未形成正式发文'],
            ['2023-03-25', '销售条线月度例会', '销售VP、销售运营、各大区经理',
             '口头传达将按新政策执行，未明确追溯范围', '会议记录'],
            ['2024-01-08', '年度销售政策宣贯会', '销售VP、HR、财务', '宣贯V3，渠道单佣金减半', '会议记录'],
            ['2024-06-20', '渠道管理专题会', '渠道总监、财务', '讨论渠道返利与销售佣金重复计算问题', '会议纪要'],
            ['2025-07-15', '半年度佣金清算会', 'CFO、销售VP、内审', '要求对2022年以来佣金开展全面清算', '会议纪要']]
FIN_RULES = [['1', '订单去重', '同一订单号在订单主表出现多次的，以首次导入记录为准'],
             ['2', '政策归属', '政策适用年度以订单发货日期为准，不以订单日期或回款日期为准'],
             ['3', '文件效力', '会议纪要、讨论稿不作为佣金计算依据，以正式发文为准'],
             ['4', '金额口径', '订单与回款金额均为含税金额，增值税率13%'],
             ['5', '退货处理', '以审批通过的退货记录为准，退款冲减回款与计提基数'],
             ['6', '坏账处理', '以财务核销单为准，坏账金额从V3净回款中扣除'],
             ['7', '统计截止', '回款与退款统计截至2025-08-31'],
             ['8', '已发口径', '佣金已发台账为财务实际支付数，含错误支付，不作为应发依据'],
             ['9', '渠道返利', '按合同约定的含税进货额计算，与销售佣金不得重复'],
             ['10', '争议裁决', '以本规则及正式政策为准，逐单核对']]
TAX_RULES = [['增值税率', '13%', '含税金额除以1.13换算为不含税'],
             ['佣金个税', '由公司代扣代缴', '不影响应发佣金口径'],
             ['返利发票', '渠道商开具6%服务费发票', '返利金额为含税口径'],
             ['汇率', '不涉及', '全部为人民币']]
MGR = [['1', '总经理', '把2022-2023年已结清的佣金按V3重算并补差', '2025-07-16'],
       ['2', '销售VP', '取消所有封顶，按理论佣金全额发放', '2025-07-16'],
       ['3', '渠道总监', '渠道返利计入销售佣金计提基数', '2025-07-17'],
       ['4', '财务经理', '按开票额计提佣金，不按回款', '2025-07-17'],
       ['5', 'HR总监', '离职人员未回款订单继续计提并发放', '2025-07-18'],
       ['6', '销售运营', '以2023-03-10会议纪要作为V2自4月1日执行的依据', '2025-07-18'],
       ['7', '客服总监', '退货订单不追回已发佣金', '2025-07-21'],
       ['8', '财务经理', '坏账核销后仍视同已回款', '2025-07-21'],
       ['9', 'HR经理', '团队池按团队人数平均发放', '2025-07-22'],
       ['10', '销售VP', '新人保底按全年12个月发放', '2025-07-22'],
       ['11', '渠道总监', '渠道单按全额3%计提', '2025-07-23'],
       ['12', '华东大区经理', '2024年起合作单仍全部归签约人', '2025-07-23']]

# ---------------- 写工作簿 ----------------
wb = openpyxl.Workbook()
wb.remove(wb.active)
sheet(wb, '销售花名册', ['工号', '姓名', '团队', '负责区域', '职级', '入职日期', '离职日期', '在职状态',
                     '月基本工资(元)', '月度目标(万元)', '备注'],
      [[r[0], r[1], r[2], r[3], r[4], r[5], r[6], '离职' if r[6] else '在职', r[8], r[9], r[10]] for r in roster],
      {1: 10, 2: 10, 3: 8, 4: 24, 5: 12, 6: 12, 7: 12, 8: 10, 9: 14, 10: 14, 11: 20})
sheet(wb, '团队目标与薪酬', ['团队', '2022年目标(万元)', '2023年目标(万元)', '2024年目标(万元)', '2025年目标(万元)', '团队负责人', '备注'],
      team_meta, {1: 8, 2: 18, 3: 18, 4: 18, 5: 18, 6: 12, 7: 26})
sheet(wb, '区域划分', ['团队', '覆盖区域', '区域代码', '备注'],
      [[t[0], t[1], f'R{i + 1}', ''] for i, t in enumerate(TEAMS)], {1: 8, 2: 40, 3: 10, 4: 20})

sheet(wb, '订单主表', ['订单号', '订单日期', '发货日期', '客户编号', '客户名称', '客户等级', '合同类型',
                    '渠道编号', '签约销售', '协作销售', '含税金额(元)', '税率', '所属团队', '备注'],
      order_rows, {1: 14, 2: 12, 3: 12, 4: 10, 5: 30, 6: 10, 7: 12, 8: 10, 9: 10, 10: 10, 11: 16, 12: 8, 13: 8, 14: 26})
sheet(wb, '回款流水', ['回款编号', '订单号', '回款日期', '回款金额(含税,元)', '回款类型', '银行流水号', '备注'],
      pay_rows, {1: 12, 2: 14, 3: 12, 4: 18, 5: 10, 6: 16, 7: 20})
sheet(wb, '退货记录', ['退货编号', '订单号', '退货日期', '退货金额(含税,元)', '退货原因', '审批状态'],
      returns, {1: 10, 2: 14, 3: 12, 4: 18, 5: 16, 6: 10})
sheet(wb, '坏账核销', ['核销编号', '订单号', '客户名称', '核销日期', '未回款金额(含税,元)', '核销依据'],
      bad, {1: 10, 2: 14, 3: 30, 4: 12, 5: 20, 6: 34})

sheet(wb, '佣金政策V1', ['条款编号', '条款', '内容'], V1, {1: 10, 2: 14, 3: 70})
sheet(wb, '佣金政策V2', ['条款编号', '条款', '内容'], V2, {1: 10, 2: 14, 3: 70})
sheet(wb, '佣金政策V3', ['条款编号', '条款', '内容'], V3, {1: 10, 2: 14, 3: 70})
sheet(wb, '政策变更通知', ['文号', '发布日期', '生效日期', '适用范围', '主要内容'],
      NOTICES, {1: 22, 2: 12, 3: 12, 4: 14, 5: 60})
sheet(wb, '经营会会议纪要', ['日期', '议题', '参会人', '结论摘要', '文件性质'],
      MEETINGS, {1: 12, 2: 24, 3: 34, 4: 44, 5: 22})

sheet(wb, '渠道商清单', ['渠道编号', '渠道名称', '所属团队', '合作起始', '合同到期', '年度保底返利(元)', '渠道经理', '备注'],
      channels, {1: 10, 2: 34, 3: 10, 4: 12, 5: 12, 6: 18, 7: 10, 8: 24})
sheet(wb, '渠道返利阶梯', ['档位', '年度进货额下限(万元)', '年度进货额上限(万元)', '返利比例', '说明'],
      [[t[0], t[1], t[2], t[3], '按年度含税进货额（发货额扣退货）计算'] for t in TIERS],
      {1: 28, 2: 20, 3: 20, 4: 10, 5: 40})
sheet(wb, '渠道协议摘要', ['渠道编号', '渠道名称', '结算年度', '返利基数', '阶梯规则', '保底条款', '结算时点'],
      [[c[0], c[1], '2022-2024', '年度含税进货额（发货额扣退货）', '见《渠道返利阶梯》',
        '年度保底30万元' if c[5] else '无', '次年1月'] for c in channels],
      {1: 10, 2: 34, 3: 12, 4: 30, 5: 18, 6: 18, 7: 10})
ch_order_rows = []
for o in all_orders:
    if o['ch']:
        ret = sum(x[3] for x in returns if x[1] == o['oid'])
        ch_order_rows.append([o['oid'], o['ch'], dstr(o['ship']), o['amt'], ret, o['amt'] - ret])
sheet(wb, '渠道订单明细', ['订单号', '渠道编号', '发货日期', '含税发货额(元)', '退货金额(元)', '净进货额(元)'],
      ch_order_rows, {1: 14, 2: 10, 3: 12, 4: 18, 5: 16, 6: 18})
sheet(wb, '渠道返利已付台账', ['渠道编号', '渠道名称', '年度', '已付返利(元)', '付款批次', '备注'],
      rebate_paid, {1: 10, 2: 34, 3: 8, 4: 16, 5: 14, 6: 20})

sheet(wb, '佣金争议台账', ['争议编号', '销售工号', '销售姓名', '订单号', '申诉日期', '申诉金额(元)', '申诉理由', '证据材料'],
      disputes, {1: 10, 2: 10, 3: 10, 4: 14, 5: 12, 6: 16, 7: 50, 8: 16})
sheet(wb, '申诉沟通记录', ['争议编号', '沟通日期', '沟通方式', '沟通对象', '沟通要点'],
      [[d[0], '2025-08-' + f'{random.randint(1, 28):02d}', random.choice(['面谈', '邮件', '电话']),
        random.choice(['销售本人', '大区经理', 'HR', '财务']),
        random.choice(['坚持按合同额计算', '同意按政策复核', '要求书面答复', '补充提交回款凭证'])] for d in disputes[:30]],
      {1: 10, 2: 12, 3: 10, 4: 12, 5: 36})

sheet(wb, '财务确认规则', ['序号', '规则', '内容'], FIN_RULES, {1: 6, 2: 12, 3: 70})
sheet(wb, '税务与换算口径', ['项目', '口径', '说明'], TAX_RULES, {1: 14, 2: 22, 3: 46})
sheet(wb, '佣金已发台账', ['工号', '姓名', '年度', '已发金额(元)', '发放批次', '备注'],
      paid_rows, {1: 10, 2: 10, 3: 8, 4: 16, 5: 16, 6: 24})
sheet(wb, '管理层要求', ['序号', '提出人', '要求内容', '提出日期'],
      MGR, {1: 6, 2: 14, 3: 60, 4: 12})

out = f'{BASE}/销售佣金清算_v1_题目.xlsx'
wb.save(out)

print('题目已生成：', out)
print('花名册', len(roster), '订单(含重复)', len(order_rows), '唯一订单', len(all_orders),
      '回款', len(pay_rows), '退货', len(returns), '坏账', len(bad),
      '渠道订单', len(ch_order_rows), '争议', len(disputes), '已发台账', len(paid_rows))
print('订单金额合计(亿)', round(sum(o['amt'] for o in all_orders) / 1e8, 2))
